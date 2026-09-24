from io import BytesIO

import pytest
from docx import Document
from pptx import Presentation

from notetaker.note_exports import document_export, readable_markdown
from test_note_edits import generated, saved_edit
from test_notes import notes, speech
from test_capture import capture
from test_workspace import setup


def extracted(response, format):
    if format == 'docx':
        return '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    if format == 'pptx':
        return '\n'.join(shape.text for slide in Presentation(BytesIO(response.content)).slides
                         for shape in slide.shapes if shape.has_text_frame)
    return response.text


@pytest.mark.parametrize('format', ['docx', 'pptx', 'txt'])
def test_saved_student_export_and_generated_history(notes, format):
    _, client, _, path, _ = notes
    revision, result, _ = saved_edit(notes)
    selected = client.get(path+'/notes').json()['editing']['selected']
    url = path+'/notes/edits/'+selected['id']+'/export'
    response = client.get(url, params={'format': format})
    assert response.status_code == 200
    text = extracted(response, format)
    assert 'My reasoning:' in text and 'lo = mid + 1' in text and 'E = mc²' in text
    assert 'Source appendix' in text
    old = client.get(path+'/notes/revisions/'+revision['id']+'/export', params={'format': format})
    assert old.status_code == 200 and 'My reasoning:' not in extracted(old, format)
    assert client.get(url, params={'format': 'exe'}).status_code == 422
    assert client.get(url.replace('/edits/', '/edits/missing'), params={'format': format}).status_code == 404
    client.cookies.clear()
    assert client.get(url, params={'format': format}).status_code == 401


def test_long_slides_keep_every_character_and_escape_xml():
    text = '# Unicode λ <script> & sources\n' + ('    x = 42; ' * 1500) + '\nEND-SOURCE-00:42'
    response = document_export('Synthetic', text, 'pptx')
    deck = Presentation(BytesIO(response.body))
    assert len(deck.slides) > 5
    bodies = [''.join(slide.shapes[1].text.split('\n')) for slide in deck.slides]
    assert ''.join(bodies) == readable_markdown(text).replace('\n', '')
    assert all(len(slide.shapes[1].text.split('\n')) <= 17 for slide in deck.slides)


@pytest.mark.parametrize('format', ['docx', 'pptx', 'txt'])
def test_final_snapshot_exports_are_frozen(notes, format):
    from test_lifecycle import finalize, reconcile_finalizations
    app, client, headers, path, _ = notes
    saved_edit(notes)
    assert finalize(client, headers, path, available=True).status_code == 202
    reconcile_finalizations(app.state.sessions)
    snapshot = client.get(path+'/finalization').json()['history'][0]['snapshot_id']
    response = client.get(path+'/final-snapshots/'+snapshot+'/export', params={'format': format})
    assert response.status_code == 200
    assert 'My reasoning:' in extracted(response, format)


def test_plain_text_does_not_decode_literal_code_entities_or_backslashes():
    assert readable_markdown('    x = "&amp;"; path = r"C:\\work"\n') == 'x = "&amp;"; path = r"C:\\work"'


def test_slides_wrap_wide_unicode_without_losing_text():
    from notetaker.note_exports import slide_lines
    text = '漢字' * 100
    lines = list(slide_lines(text))
    assert ''.join(lines) == text
    assert max(map(len, lines)) <= 43

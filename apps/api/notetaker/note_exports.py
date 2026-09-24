"""Portable saved-note exports. No model calls or external resource loading."""
from io import BytesIO
import re
import unicodedata

from fastapi.responses import Response


def readable_markdown(text: str) -> str:
    from markdown_it import MarkdownIt
    parts = []
    heading = ''
    for token in MarkdownIt('commonmark', {'html': False}).parse(text):
        if token.type == 'heading_open':
            heading = '#' * int(token.tag[1]) + ' '
        elif token.type == 'inline':
            body = ''.join('\n' if child.type in ('softbreak', 'hardbreak') else
                           child.content if child.type in ('text', 'code_inline') else ''
                           for child in token.children or [])
            parts.append(heading + body)
            heading = ''
        elif token.type in ('code_block', 'fence'):
            parts.append(token.content.rstrip('\n'))
    return '\n\n'.join(parts)


def slide_lines(text: str):
    for line in text.split('\n'):
        chunk, columns = '', 0
        for character in line:
            width = (4 - columns % 4) if character == '\t' else (0 if unicodedata.combining(character)
                    else 2 if unicodedata.east_asian_width(character) in ('W', 'F') else 1)
            if columns + width > 86 and chunk:
                yield chunk
                chunk, columns = '', 0
            chunk += character
            columns += width
        yield chunk


def document_export(title: str, text: str, format: str) -> Response:
    text = readable_markdown(text)
    # XML 1.0 forbids control characters; tabs and line breaks remain intact.
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)
    title = re.sub(r'[\x00-\x1f]', ' ', title)
    stream = BytesIO()
    if format == 'txt':
        data = text.encode('utf-8')
        media = 'text/plain; charset=utf-8'
    elif format == 'docx':
        from docx import Document
        from docx.shared import Pt
        document = Document()
        document.core_properties.title = title
        document.styles['Normal'].font.name = 'Calibri'
        document.styles['Normal'].font.size = Pt(11)
        fenced = False
        for line in text.split('\n'):
            if line.startswith('```'):
                fenced = not fenced
                document.add_paragraph(line).runs[0].font.name = 'Consolas'
                continue
            heading = re.match(r'^(#{1,6}) (.*)$', line) if not fenced else None
            if heading:
                document.add_heading(heading[2], level=min(4, len(heading[1])))
            else:
                paragraph = document.add_paragraph(line)
                if fenced and paragraph.runs:
                    paragraph.runs[0].font.name = 'Consolas'
        document.save(stream)
        data = stream.getvalue()
        media = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif format == 'pptx':
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        presentation.core_properties.title = title
        # Fixed-width reading deck: bound both columns and rows. Never truncate sources.
        lines = list(slide_lines(text))
        pages = [lines[i:i+17] for i in range(0, len(lines), 17)] or [[]]
        for index, page in enumerate(pages):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string('F2F3F1')
            heading = slide.shapes.add_textbox(Inches(.65), Inches(.3), Inches(12), Inches(.6)).text_frame
            heading.text = f'Lecture notes · {index+1} / {len(pages)}'
            heading.paragraphs[0].font.size = Pt(24)
            box = slide.shapes.add_textbox(Inches(.65), Inches(1.1), Inches(12), Inches(5.9)).text_frame
            box.word_wrap = False
            for n, line in enumerate(page):
                paragraph = box.paragraphs[0] if n == 0 else box.add_paragraph()
                paragraph.text = line
                paragraph.font.name = 'Consolas'
                paragraph.font.size = Pt(16)
                paragraph.font.color.rgb = RGBColor.from_string('1B263B')
                paragraph.space_after = Pt(3)
            slide.notes_slide.notes_text_frame.text = '\n'.join(page)
        presentation.save(stream)
        data = stream.getvalue()
        media = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
    else:
        raise ValueError('Unsupported document format')
    return Response(data, media_type=media, headers={
        'Content-Disposition': f'attachment; filename="lecture-notes.{format}"',
        'X-Content-Type-Options': 'nosniff',
    })
